import io

import docx
import pytest

from app.services.parsing import (
    UnsupportedFileType,
    is_supported,
    parse_document,
)


def test_parse_docx_roundtrip():
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    buf = io.BytesIO()
    document.save(buf)

    segments = parse_document("report.docx", buf.getvalue())
    assert len(segments) == 1
    assert "First paragraph." in segments[0].text
    assert "Second paragraph." in segments[0].text


def test_parse_txt_returns_single_segment():
    segments = parse_document("notes.txt", b"hello world")
    assert len(segments) == 1
    assert segments[0].text == "hello world"
    assert segments[0].page is None


def test_parse_markdown():
    segments = parse_document("readme.md", b"# Title\n\nBody text")
    assert len(segments) == 1
    assert "Title" in segments[0].text


def test_empty_file_yields_no_segments():
    assert parse_document("empty.txt", b"   ") == []


def test_nul_bytes_are_stripped():
    segments = parse_document("weird.txt", b"hello\x00world\x00")
    assert len(segments) == 1
    assert "\x00" not in segments[0].text
    assert segments[0].text == "helloworld"


def test_unsupported_extension_raises():
    with pytest.raises(UnsupportedFileType):
        parse_document("image.png", b"\x89PNG")


def test_no_extension_defaults_to_text():
    segments = parse_document("LICENSE", b"MIT")
    assert segments[0].text == "MIT"


@pytest.mark.parametrize(
    "name,expected",
    [("a.pdf", True), ("a.md", True), ("a.txt", True), ("a.docx", True), ("a.png", False), ("a", False)],
)
def test_is_supported(name, expected):
    assert is_supported(name) is expected
